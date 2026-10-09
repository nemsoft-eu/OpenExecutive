"""Unit tests for the Executive Artifacts HTTP surface.

The route handlers are exercised DIRECTLY (not via `create_app`/TestClient)
because app construction initializes the ChromaDB knowledge store, which the
artifacts endpoints don't touch. Both underlying stores resolve their
module-level DB_PATH at call time, so we point them at one temp SQLite file
and seed it through the real store APIs.
"""
from __future__ import annotations

import io
import json
from pathlib import Path
from types import SimpleNamespace
from typing import Any

import pytest
from docx import Document
from fastapi import HTTPException
from openpyxl import load_workbook

from openexecutive.alerts import store as alerts_store
from openexecutive.api.routes import artifacts as artifacts_route
from openexecutive.orchestrator.artifact_records import Viewer
from openexecutive.workflows import persistence as wf_persistence


@pytest.fixture()
def db(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> Path:
    """One SQLite file holding both the alerts and workflow_runs tables.

    The route calls the store helpers WITHOUT a db_path, so they read each
    module's module-level DB_PATH — patch both to the temp file.
    """
    db_path = tmp_path / "episodic.db"
    alerts_store.initialize_db(db_path)
    wf_persistence.initialize_runs_db(db_path)
    monkeypatch.setattr(alerts_store, "DB_PATH", db_path)
    monkeypatch.setattr(wf_persistence, "DB_PATH", db_path)
    # The principal is calling unless a test says otherwise: a draft with no
    # owner is theirs.
    monkeypatch.setattr(artifacts_route, "_viewer", lambda request: PRINCIPAL)
    return db_path


# Stands in for the request; `_viewer` (patched above) decides who it is.
_REQ: Any = object()
PRINCIPAL = Viewer(person_id=1, is_principal=True)


def _seed_draft(db: Path, external_id: str, headline: str) -> int:
    aid = alerts_store.insert_alert(
        source="artifact", external_id=external_id, severity="medium",
        headline=headline, body=f"## {headline}\n\nThe body of {headline}.",
        suggested_action=f"Why {headline} matters.", topic_tags=["artifact"],
        db_path=db,
    )
    assert aid is not None
    return aid


def _seed_run(db: Path, run_id: str, title: str) -> None:
    wf_persistence.create_run(run_id, "board_prep", title, {}, db_path=db)
    wf_persistence.complete_run(run_id, f"# {title}\n\nDeck body.", db_path=db)


async def test_list_merges_both_sources_sorted_desc(db: Path) -> None:
    aid = _seed_draft(db, "a-1", "Competitor teardown")
    _seed_run(db, "run-1", "Q2 Board Deck")
    # A non-artifact alert and an unfinished run must NOT appear.
    alerts_store.insert_alert(
        source="email", external_id="e-1", severity="high",
        headline="Inbound", body="x", db_path=db,
    )
    wf_persistence.create_run("run-wip", "board_prep", "WIP", {}, db_path=db)

    result = await artifacts_route.list_artifacts(_REQ)
    items = result["artifacts"]

    by_id = {a.id: a for a in items}
    assert set(by_id) == {f"alert:{aid}", "run:run-1"}
    assert by_id[f"alert:{aid}"].kind == "draft"
    assert by_id[f"alert:{aid}"].source_label == "Drafted by Executive"
    assert by_id[f"alert:{aid}"].preview  # drafts carry a preview
    assert by_id["run:run-1"].kind == "workflow"
    assert by_id["run:run-1"].source_label == "board_prep"
    assert by_id["run:run-1"].preview is None  # runs don't, by design

    created = [a.created_at for a in items]
    assert created == sorted(created, reverse=True)


async def test_list_includes_acked_draft(db: Path) -> None:
    aid = _seed_draft(db, "a-ack", "Acked memo")
    alerts_store.set_status(aid, "ack", db_path=db)
    ids = {a.id for a in (await artifacts_route.list_artifacts(_REQ))["artifacts"]}
    assert f"alert:{aid}" in ids


async def test_detail_draft_returns_body_and_rationale(db: Path) -> None:
    aid = _seed_draft(db, "a-d", "Memo")
    detail = await artifacts_route.get_artifact(f"alert:{aid}", _REQ)
    assert detail.kind == "draft"
    assert "The body of Memo." in detail.body
    assert detail.rationale == "Why Memo matters."


async def test_detail_run_returns_artifact_body(db: Path) -> None:
    _seed_run(db, "run-x", "Deck")
    detail = await artifacts_route.get_artifact("run:run-x", _REQ)
    assert detail.kind == "workflow"
    assert "Deck body." in detail.body
    assert detail.rationale is None


async def test_detail_404_for_non_artifact_alert(db: Path) -> None:
    aid = alerts_store.insert_alert(
        source="email", external_id="e-2", severity="high",
        headline="Inbound", body="x", db_path=db,
    )
    with pytest.raises(HTTPException) as exc:
        await artifacts_route.get_artifact(f"alert:{aid}", _REQ)
    assert exc.value.status_code == 404


async def test_detail_404_for_unknown_ids(db: Path) -> None:
    for bad in ("alert:9999", "run:nope"):
        with pytest.raises(HTTPException) as exc:
            await artifacts_route.get_artifact(bad, _REQ)
        assert exc.value.status_code == 404


async def test_detail_404_for_incomplete_run(db: Path) -> None:
    wf_persistence.create_run("run-wip2", "board_prep", "WIP", {}, db_path=db)
    with pytest.raises(HTTPException) as exc:
        await artifacts_route.get_artifact("run:run-wip2", _REQ)
    assert exc.value.status_code == 404


@pytest.mark.parametrize("bad_id", ["alert:notanint", "garbage", "weird:1"])
async def test_detail_400_for_malformed_id(db: Path, bad_id: str) -> None:
    with pytest.raises(HTTPException) as exc:
        await artifacts_route.get_artifact(bad_id, _REQ)
    assert exc.value.status_code == 400


# ---------------------------------------------------------------------------
# Archive / restore / delete
# ---------------------------------------------------------------------------


async def _active_ids(db: Path) -> set[str]:
    return {a.id for a in (await artifacts_route.list_artifacts(_REQ))["artifacts"]}


async def _archived_ids(db: Path) -> set[str]:
    result = await artifacts_route.list_artifacts(_REQ, archived=True)
    return {a.id for a in result["artifacts"]}


async def test_archive_then_restore_draft(db: Path) -> None:
    aid = _seed_draft(db, "a-arch", "Archivable memo")
    cid = f"alert:{aid}"

    assert cid in await _active_ids(db)
    assert cid not in await _archived_ids(db)

    res = await artifacts_route.archive_artifact(cid, _REQ)
    assert res == {"status": "archived", "id": cid}
    # Gone from active, present in archived — a clean swap, not a superset.
    assert cid not in await _active_ids(db)
    assert cid in await _archived_ids(db)
    # Detail still resolves an archived artifact (so the gallery can open it).
    assert (await artifacts_route.get_artifact(cid, _REQ)).id == cid

    res = await artifacts_route.restore_artifact(cid, _REQ)
    assert res == {"status": "restored", "id": cid}
    assert cid in await _active_ids(db)
    assert cid not in await _archived_ids(db)


async def test_archived_at_surfaces_in_responses(db: Path) -> None:
    """The detail page's Archive/Restore toggle depends on `archived_at` being
    populated in both list and detail responses (for both kinds)."""
    aid = _seed_draft(db, "a-flag", "Flagged memo")
    _seed_run(db, "run-flag", "Flagged deck")
    cid_alert, cid_run = f"alert:{aid}", "run:run-flag"

    # Active: archived_at is None everywhere.
    for item in (await artifacts_route.list_artifacts(_REQ))["artifacts"]:
        assert item.archived_at is None
    assert (await artifacts_route.get_artifact(cid_alert, _REQ)).archived_at is None
    assert (await artifacts_route.get_artifact(cid_run, _REQ)).archived_at is None

    await artifacts_route.archive_artifact(cid_alert, _REQ)
    await artifacts_route.archive_artifact(cid_run, _REQ)

    # Archived: archived_at is a real ISO timestamp in list + detail, both kinds.
    archived = {a.id: a for a in (await artifacts_route.list_artifacts(_REQ, archived=True))["artifacts"]}
    assert archived[cid_alert].archived_at
    assert archived[cid_run].archived_at
    assert (await artifacts_route.get_artifact(cid_alert, _REQ)).archived_at
    assert (await artifacts_route.get_artifact(cid_run, _REQ)).archived_at


async def test_archive_then_restore_run(db: Path) -> None:
    _seed_run(db, "run-arch", "Archivable deck")
    cid = "run:run-arch"

    assert cid in await _active_ids(db)
    await artifacts_route.archive_artifact(cid, _REQ)
    assert cid not in await _active_ids(db)
    assert cid in await _archived_ids(db)

    await artifacts_route.restore_artifact(cid, _REQ)
    assert cid in await _active_ids(db)
    assert cid not in await _archived_ids(db)


async def test_delete_draft_removes_it(db: Path) -> None:
    aid = _seed_draft(db, "a-del", "Deletable memo")
    cid = f"alert:{aid}"
    res = await artifacts_route.delete_artifact(cid, _REQ)
    assert res == {"status": "deleted", "id": cid}
    assert cid not in await _active_ids(db)
    with pytest.raises(HTTPException) as exc:
        await artifacts_route.get_artifact(cid, _REQ)
    assert exc.value.status_code == 404


async def test_delete_run_removes_it(db: Path) -> None:
    _seed_run(db, "run-del", "Deletable deck")
    cid = "run:run-del"
    await artifacts_route.delete_artifact(cid, _REQ)
    assert cid not in await _active_ids(db)
    with pytest.raises(HTTPException) as exc:
        await artifacts_route.get_artifact(cid, _REQ)
    assert exc.value.status_code == 404


@pytest.mark.parametrize(
    "mutator",
    [
        artifacts_route.archive_artifact,
        artifacts_route.restore_artifact,
        artifacts_route.delete_artifact,
    ],
)
async def test_mutators_404_for_unknown_ids(db: Path, mutator) -> None:  # type: ignore[no-untyped-def]
    for bad in ("alert:9999", "run:nope"):
        with pytest.raises(HTTPException) as exc:
            await mutator(bad, _REQ)
        assert exc.value.status_code == 404


@pytest.mark.parametrize(
    "mutator",
    [
        artifacts_route.archive_artifact,
        artifacts_route.restore_artifact,
        artifacts_route.delete_artifact,
    ],
)
async def test_mutators_404_for_non_artifact_alert(db: Path, mutator) -> None:  # type: ignore[no-untyped-def]
    # A real alert row, but source='email' — the artifact routes must not
    # become a general alert mutator.
    aid = alerts_store.insert_alert(
        source="email", external_id="e-mut", severity="high",
        headline="Inbound", body="x", db_path=db,
    )
    with pytest.raises(HTTPException) as exc:
        await mutator(f"alert:{aid}", _REQ)
    assert exc.value.status_code == 404


@pytest.mark.parametrize(
    "mutator",
    [
        artifacts_route.archive_artifact,
        artifacts_route.restore_artifact,
        artifacts_route.delete_artifact,
    ],
)
@pytest.mark.parametrize("bad_id", ["alert:notanint", "garbage", "weird:1"])
async def test_mutators_400_for_malformed_id(
    db: Path, mutator, bad_id: str  # type: ignore[no-untyped-def]
) -> None:
    with pytest.raises(HTTPException) as exc:
        await mutator(bad_id, _REQ)
    assert exc.value.status_code == 400


def test_store_set_alert_archived_round_trip(db: Path) -> None:
    aid = _seed_draft(db, "a-store", "Store memo")
    assert [a.id for a in alerts_store.list_artifact_alerts(owner_person_id=None, include_unowned=True, db_path=db)] == [aid]
    assert alerts_store.list_artifact_alerts(owner_person_id=None, include_unowned=True, db_path=db, archived=True) == []

    assert alerts_store.set_alert_archived(aid, True, db_path=db) is True
    assert alerts_store.list_artifact_alerts(owner_person_id=None, include_unowned=True, db_path=db) == []
    assert [a.id for a in alerts_store.list_artifact_alerts(owner_person_id=None, include_unowned=True, db_path=db, archived=True)] == [aid]

    assert alerts_store.set_alert_archived(aid, False, db_path=db) is True
    assert [a.id for a in alerts_store.list_artifact_alerts(owner_person_id=None, include_unowned=True, db_path=db)] == [aid]
    # Unknown id reports no row updated.
    assert alerts_store.set_alert_archived(99999, True, db_path=db) is False


def test_store_set_run_archived_round_trip(db: Path) -> None:
    _seed_run(db, "run-store", "Store deck")
    assert [r["run_id"] for r in wf_persistence.list_artifact_runs(visible_to=None, db_path=db)] == ["run-store"]
    assert wf_persistence.list_artifact_runs(visible_to=None, db_path=db, archived=True) == []

    assert wf_persistence.set_run_archived("run-store", True, db_path=db) is True
    assert wf_persistence.list_artifact_runs(visible_to=None, db_path=db) == []
    assert [r["run_id"] for r in wf_persistence.list_artifact_runs(visible_to=None, db_path=db, archived=True)] == ["run-store"]

    assert wf_persistence.set_run_archived("run-store", False, db_path=db) is True
    assert [r["run_id"] for r in wf_persistence.list_artifact_runs(visible_to=None, db_path=db)] == ["run-store"]
    assert wf_persistence.set_run_archived("nope", True, db_path=db) is False


# --------------------------------------------------------------------------- #
# Formats and downloads
# --------------------------------------------------------------------------- #




def _seed_format(db: Path, external_id: str, fmt: str, body: str, **extra: object) -> str:
    aid = alerts_store.insert_alert(
        source="artifact", external_id=external_id, severity="medium",
        headline=f"{fmt} artifact", body=body, suggested_action="why",
        topic_tags=["artifact"], artifact_format=fmt, db_path=db, **extra,  # type: ignore[arg-type]
    )
    assert aid is not None
    return f"alert:{aid}"


async def test_list_carries_format_fields(db: Path) -> None:
    md = _seed_format(db, "m", "markdown", "# T\n\nHello")
    link = _seed_format(db, "l", "link", "Summary", artifact_url="https://x.example/s",
                        artifact_link_label="Notion page")
    _seed_run(db, "run-f", "Deck")

    by_id = {a.id: a for a in (await artifacts_route.list_artifacts(_REQ))["artifacts"]}
    assert by_id[md].format == "markdown"
    assert by_id[md].downloads == ["markdown", "docx"]
    assert by_id[link].format == "link"
    assert by_id[link].format_label == "Notion page"
    assert by_id[link].downloads == []
    assert by_id[link].external_url == "https://x.example/s"
    # Workflow output is Markdown and can be exported to Word.
    assert by_id["run:run-f"].format == "markdown"
    assert by_id["run:run-f"].downloads == ["markdown", "docx"]


async def test_detail_body_per_format(db: Path) -> None:
    html = _seed_format(db, "h", "html", "<h1>Hi</h1><p>Text</p>")
    sheet = _seed_format(db, "x", "xlsx", json.dumps({
        "summary": "S", "sheets": [{"name": "A", "columns": ["c"], "rows": [[1]]}],
    }))
    html_detail = await artifacts_route.get_artifact(html, _REQ)
    assert html_detail.body == "<h1>Hi</h1><p>Text</p>"  # raw, for the sandboxed iframe
    assert html_detail.preview == "Hi Text"
    sheet_detail = await artifacts_route.get_artifact(sheet, _REQ)
    assert "| c |" in sheet_detail.body


async def _download_headers(resp) -> dict[str, str]:  # type: ignore[no-untyped-def]
    return {k.lower(): v for k, v in resp.headers.items()}


async def test_download_markdown_and_docx_export(db: Path) -> None:
    aid = _seed_draft(db, "d-1", "Board Memo: Q3!")
    resp = await artifacts_route.download_artifact(f"alert:{aid}", _REQ)
    headers = await _download_headers(resp)
    assert resp.body.startswith(b"## Board Memo")
    assert headers["content-type"].startswith("text/markdown")
    assert headers["content-disposition"] == 'attachment; filename="board-memo-q3.md"'
    assert headers["x-content-type-options"] == "nosniff"
    assert headers["content-security-policy"] == "sandbox"

    word = await artifacts_route.download_artifact(f"alert:{aid}", _REQ, as_="docx")
    assert (await _download_headers(word))["content-disposition"].endswith('.docx"')
    doc = Document(io.BytesIO(word.body))
    assert any(p.text == "Board Memo: Q3!" for p in doc.paragraphs)


async def test_download_workflow_run_as_docx(db: Path) -> None:
    _seed_run(db, "run-d", "Q2 Board Deck")
    resp = await artifacts_route.download_artifact("run:run-d", _REQ, as_="docx")
    doc = Document(io.BytesIO(resp.body))
    assert any(p.text == "Deck body." for p in doc.paragraphs)


async def test_download_xlsx_renders_workbook(db: Path) -> None:
    cid = _seed_format(db, "x2", "xlsx", json.dumps({
        "summary": "", "sheets": [{"name": "Data", "columns": ["a", "b"], "rows": [[1, "two"]]}],
    }))
    resp = await artifacts_route.download_artifact(cid, _REQ)
    wb = load_workbook(io.BytesIO(resp.body))
    assert wb["Data"]["B2"].value == "two"


async def test_download_html_is_attachment(db: Path) -> None:
    cid = _seed_format(db, "h2", "html", "<p>x</p>")
    headers = await _download_headers(await artifacts_route.download_artifact(cid, _REQ))
    assert headers["content-disposition"].startswith("attachment;")
    assert headers["content-disposition"].endswith('.html"')
    assert headers["x-content-type-options"] == "nosniff"


async def test_download_404_for_link_and_unsupported_target(db: Path) -> None:
    link = _seed_format(db, "l2", "link", "S", artifact_url="https://x.example")
    with pytest.raises(HTTPException) as exc:
        await artifacts_route.download_artifact(link, _REQ)
    assert exc.value.status_code == 404
    html = _seed_format(db, "h3", "html", "<p>x</p>")
    with pytest.raises(HTTPException) as exc:
        await artifacts_route.download_artifact(html, _REQ, as_="xlsx")
    assert exc.value.status_code == 404


async def test_download_404_for_non_artifact_alert(db: Path) -> None:
    other = alerts_store.insert_alert(
        source="email", external_id="e-dl", severity="high",
        headline="Inbound", body="secret", db_path=db,
    )
    with pytest.raises(HTTPException) as exc:
        await artifacts_route.download_artifact(f"alert:{other}", _REQ)
    assert exc.value.status_code == 404


async def test_delete_unindexes(db: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    removed: list[str] = []

    async def _fake(artifact_id: str) -> None:
        removed.append(artifact_id)

    monkeypatch.setattr(
        "openexecutive.orchestrator.artifact_tools.unindex_artifact", _fake
    )
    aid = _seed_draft(db, "d-del", "Gone")
    # A zero-padded spelling still unindexes the canonical id.
    await artifacts_route.delete_artifact(f"alert:0{aid}", _REQ)
    assert removed == [f"alert:{aid}"]


async def test_archive_unindexes_and_restore_reindexes(
    db: Path, monkeypatch: pytest.MonkeyPatch,
) -> None:
    calls: list[tuple] = []

    async def _unindex(artifact_id: str) -> None:
        calls.append(("unindex", artifact_id))

    async def _index(
        artifact_id: str, title: str, fmt: str, stored: str, owner: int | None
    ) -> None:
        calls.append(("index", artifact_id, title, fmt))

    monkeypatch.setattr("openexecutive.orchestrator.artifact_tools.unindex_artifact", _unindex)
    monkeypatch.setattr("openexecutive.orchestrator.artifact_tools.index_artifact", _index)
    aid = _seed_draft(db, "d-arch", "Memo")
    await artifacts_route.archive_artifact(f"alert:{aid}", _REQ)
    await artifacts_route.restore_artifact(f"alert:{aid}", _REQ)
    assert calls == [
        ("unindex", f"alert:{aid}"),
        ("index", f"alert:{aid}", "Memo", "markdown"),
    ]


async def test_filename_falls_back_to_id(db: Path) -> None:
    aid = _seed_draft(db, "d-sym", "!!!")
    resp = await artifacts_route.download_artifact(f"alert:{aid}", _REQ)
    headers = await _download_headers(resp)
    assert headers["content-disposition"] == f'attachment; filename="alert-{aid}.md"'


# --------------------------------------------------------------------------- #
# Ownership: every route answers only for the caller's own documents
# --------------------------------------------------------------------------- #


async def test_routes_show_each_caller_only_their_own(
    db: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    sams = alerts_store.insert_alert(
        source="artifact", external_id="sam-1", severity="medium",
        headline="Sam's notes", body="private to Sam", topic_tags=["artifact"],
        owner_person_id=7, db_path=db,
    )
    principals = _seed_draft(db, "p-1", "Board memo")  # no owner: the principal's
    _seed_run(db, "team-run", "Team deck")              # scheduled: the team's
    wf_persistence.create_run(
        "sams-run", "board_prep", "Sam's deck", {}, db_path=db, owner_person_id=7
    )
    wf_persistence.complete_run("sams-run", "# Sam", db_path=db)
    sam = Viewer(person_id=7)

    def listing() -> Any:
        return artifacts_route.list_artifacts(_REQ)

    monkeypatch.setattr(artifacts_route, "_viewer", lambda request: sam)
    assert {a.title for a in (await listing())["artifacts"]} == {
        "Sam's notes", "Sam's deck", "Team deck",
    }
    assert (await artifacts_route.get_artifact(f"alert:{sams}", _REQ)).body == "private to Sam"

    # The principal, a teammate and someone off the roster get a 404 for
    # every route, exactly as for a missing id.
    for viewer in (PRINCIPAL, Viewer(person_id=8), Viewer(person_id=None)):
        monkeypatch.setattr(artifacts_route, "_viewer", lambda request, v=viewer: v)
        titles = {a.title for a in (await listing())["artifacts"]}
        assert "Sam's notes" not in titles and "Sam's deck" not in titles
        assert "Team deck" in titles
        for cid in (f"alert:{sams}", "run:sams-run"):
            for call in (
                artifacts_route.get_artifact,
                artifacts_route.download_artifact,
                artifacts_route.archive_artifact,
                artifacts_route.restore_artifact,
                artifacts_route.delete_artifact,
            ):
                with pytest.raises(HTTPException) as exc:
                    await call(cid, _REQ)
                assert exc.value.status_code == 404
    # Only the principal sees a draft with no owner.
    monkeypatch.setattr(artifacts_route, "_viewer", lambda request: sam)
    with pytest.raises(HTTPException):
        await artifacts_route.get_artifact(f"alert:{principals}", _REQ)
    # Nothing was changed by the refused calls.
    alert = alerts_store.get_alert(sams, db_path=db)
    assert alert is not None and alert.archived_at is None
    assert wf_persistence.get_run("sams-run", db_path=db) is not None


# --------------------------------------------------------------------------- #
# Extension collections (orchestrator/extensions.py)
# --------------------------------------------------------------------------- #


@pytest.fixture()
def widgets(monkeypatch: pytest.MonkeyPatch) -> Any:
    from openexecutive.orchestrator import extensions

    extensions._reset_for_tests()
    monkeypatch.setattr(extensions, "_loaded", True)
    seen: list[tuple[str, str]] = []
    refuse: set[str] = set()

    async def on_change(artifact_id: str, change: str) -> None:
        if change in refuse:
            raise RuntimeError("no")
        seen.append((artifact_id, change))

    extensions.register_collection(
        extensions.Collection(name="widgets", label="Widgets", on_change=on_change)
    )

    async def _noop(*_a: Any) -> None:
        return None

    monkeypatch.setattr("openexecutive.orchestrator.artifact_tools.unindex_artifact", _noop)
    monkeypatch.setattr("openexecutive.orchestrator.artifact_tools.index_artifact", _noop)
    yield SimpleNamespace(seen=seen, refuse=refuse)
    extensions._reset_for_tests()


def _seed_widget(db: Path) -> int:
    aid = alerts_store.insert_alert(
        source="artifact", external_id="w-1", severity="medium",
        headline="Hiring board", body="A board.", topic_tags=["artifact", "collection:widgets"],
        artifact_format="link", artifact_url="https://w.example/1", artifact_link_label="Widget",
        db_path=db,
    )
    assert aid is not None
    return aid


async def test_a_collection_document_carries_its_collection(db: Path, widgets: Any) -> None:
    aid = _seed_widget(db)
    plain = _seed_draft(db, "d-plain", "Memo")
    by_id = {a.id: a for a in (await artifacts_route.list_artifacts(_REQ))["artifacts"]}
    assert (by_id[f"alert:{aid}"].collection, by_id[f"alert:{aid}"].collection_label) == ("widgets", "Widgets")
    assert by_id[f"alert:{plain}"].collection is None


async def test_an_unregistered_collection_tag_is_a_plain_draft(db: Path) -> None:
    from openexecutive.orchestrator import extensions

    extensions._reset_for_tests()
    aid = _seed_widget(db)
    detail = await artifacts_route.get_artifact(f"alert:{aid}", _REQ)
    assert detail.collection is None and detail.collection_label is None


async def test_the_collection_hears_each_change_first(db: Path, widgets: Any) -> None:
    aid = _seed_widget(db)
    await artifacts_route.archive_artifact(f"alert:{aid}", _REQ)
    await artifacts_route.restore_artifact(f"alert:{aid}", _REQ)
    await artifacts_route.delete_artifact(f"alert:{aid}", _REQ)
    assert widgets.seen == [
        (f"alert:{aid}", "archived"), (f"alert:{aid}", "restored"), (f"alert:{aid}", "deleted"),
    ]


async def test_a_refused_change_does_not_happen(db: Path, widgets: Any) -> None:
    aid = _seed_widget(db)
    widgets.refuse.update({"archived", "deleted"})
    for mutate in (artifacts_route.archive_artifact, artifacts_route.delete_artifact):
        with pytest.raises(HTTPException) as exc:
            await mutate(f"alert:{aid}", _REQ)
        assert exc.value.status_code == 502
    detail = await artifacts_route.get_artifact(f"alert:{aid}", _REQ)
    assert detail.archived_at is None


async def test_a_change_that_changes_nothing_is_not_reported(db: Path, widgets: Any) -> None:
    aid = _seed_widget(db)
    await artifacts_route.restore_artifact(f"alert:{aid}", _REQ)
    await artifacts_route.archive_artifact(f"alert:{aid}", _REQ)
    await artifacts_route.archive_artifact(f"alert:{aid}", _REQ)
    assert widgets.seen == [(f"alert:{aid}", "archived")]
